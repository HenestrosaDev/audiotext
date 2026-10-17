import json
from pathlib import Path

import pytest
from docx import Document

from models.config.config_subtitles import ConfigSubtitles
from models.summary import Chapter, TranscriptSummary
from models.transcript_segment import TranscriptSegment, TranscriptWord
from utils.config_manager import ConfigManager
from utils.exporters import (
    SUMMARY_EXPORTERS,
    ExportDocument,
    Paragraph,
    available_formats,
    export,
    export_collection,
    export_summary,
    paragraphs,
)

SEGMENTS = [
    TranscriptSegment(0.0, 1.5, "Hello."),
    TranscriptSegment(3661.25, 3662.0, "Bye.", speaker="SPEAKER_01"),
]
TEXT = "Hello. Bye."
SUMMARY = TranscriptSummary(
    summary="A greeting.",
    key_points=("They say hello",),
    chapters=(Chapter(0.0, "Greeting"), Chapter(3661.0, "Farewell")),
)


def document(**kwargs: object) -> ExportDocument:
    values: dict[str, object] = {"title": "Talk", "text": TEXT, "segments": SEGMENTS}
    return ExportDocument(**(values | kwargs))  # type: ignore[arg-type]


def test_srt(tmp_path: Path) -> None:
    path = export(tmp_path / "talk.mp4", "srt", document())

    assert path == tmp_path / "talk.srt"
    assert path.read_text(encoding="utf-8") == (
        "1\n00:00:00,000 --> 00:00:01,500\nHello.\n\n"
        "2\n01:01:01,250 --> 01:01:02,000\n[SPEAKER_01]: Bye.\n"
    )


def ass_events(path: Path) -> list[str]:
    return [
        line
        for line in path.read_text(encoding="utf-8").splitlines()
        if line.startswith("Dialogue:")
    ]


def test_ass(tmp_path: Path) -> None:
    path = export(tmp_path / "talk.mp4", "ass", document())

    assert path == tmp_path / "talk.ass"
    content = path.read_text(encoding="utf-8")
    assert content.startswith("[Script Info]\nScriptType: v4.00+\n")
    assert "PlayResX: 1920\nPlayResY: 1080\n" in content
    assert ass_events(path) == [
        "Dialogue: 0,0:00:00.00,0:00:01.50,Default,,0,0,0,,Hello.",
        "Dialogue: 0,1:01:01.25,1:01:02.00,Speaker1,SPEAKER_01,0,0,0,,"
        "[SPEAKER_01]: Bye.",
    ]


def test_ass_gives_each_speaker_a_color(tmp_path: Path) -> None:
    segments = [
        TranscriptSegment(0, 1, "Hi.", speaker="SPEAKER_00"),
        TranscriptSegment(1, 2, "Hello.", speaker="SPEAKER_01"),
        TranscriptSegment(2, 3, "Bye.", speaker="SPEAKER_00"),
    ]

    content = export(tmp_path / "talk", "ass", document(segments=segments))
    styles = [
        line.split(",")[:4]
        for line in content.read_text(encoding="utf-8").splitlines()
        if line.startswith("Style:")
    ]

    assert [style[0] for style in styles] == [
        "Style: Default",
        "Style: Speaker1",
        "Style: Speaker2",
    ]
    # Each speaker has a color of its own
    assert len({style[3] for style in styles}) == 3
    assert [event.split(",")[3] for event in ass_events(content)] == [
        "Speaker1",
        "Speaker2",
        "Speaker1",
    ]


def test_ass_escapes_the_tags_and_breaks_the_lines(tmp_path: Path) -> None:
    section = ConfigSubtitles.Key.SECTION
    ConfigManager.modify_value(section, ConfigSubtitles.Key.MAX_LINE_WIDTH, "12")
    ConfigManager.modify_value(section, ConfigSubtitles.Key.MAX_LINE_COUNT, "2")
    segment = TranscriptSegment(0, 2, r"Say {hi} to C:\Users now")

    path = export(tmp_path / "talk", "ass", document(segments=[segment]))

    assert ass_events(path) == [
        r"Dialogue: 0,0:00:00.00,0:00:02.00,Default,,0,0,0,,Say (hi) to\NC:/Users now"
    ]


def test_ass_of_a_vertical_video(tmp_path: Path) -> None:
    content = export(tmp_path / "talk", "ass", document(is_vertical=True)).read_text(
        encoding="utf-8"
    )

    assert "PlayResX: 1080\nPlayResY: 1920\n" in content


def test_srt_splits_long_segments_as_configured(tmp_path: Path) -> None:
    section = ConfigSubtitles.Key.SECTION
    ConfigManager.modify_value(section, ConfigSubtitles.Key.MAX_LINE_WIDTH, "10")
    ConfigManager.modify_value(section, ConfigSubtitles.Key.MAX_LINE_COUNT, "2")
    words = tuple(
        TranscriptWord(float(idx), idx + 0.5, word)
        for idx, word in enumerate(
            ["I", "just", "heard", "it", "drop", "over", "there"]
        )
    )
    segment = TranscriptSegment(
        0.0, 7.0, "I just heard it drop over there", "SPEAKER_00", words
    )

    content = export(tmp_path / "talk", "srt", document(segments=[segment]))

    # The speaker doesn't count in the width, and starts the first subtitle
    assert content.read_text(encoding="utf-8") == (
        "1\n00:00:00,000 --> 00:00:04,000\n[SPEAKER_00]: I just\nheard it\n\n"
        "2\n00:00:04,000 --> 00:00:07,000\ndrop over\nthere\n"
    )


def test_vertical_videos_and_wide_scripts_have_narrower_subtitles(
    tmp_path: Path,
) -> None:
    section = ConfigSubtitles.Key.SECTION
    ConfigManager.modify_value(section, ConfigSubtitles.Key.MAX_LINE_COUNT, "1")
    ConfigManager.modify_value(
        section, ConfigSubtitles.Key.MAX_LINE_COUNT_VERTICAL, "1"
    )
    ConfigManager.modify_value(section, ConfigSubtitles.Key.MAX_LINE_WIDTH_CJK, "4")
    latin = [TranscriptSegment(0.0, 2.0, "Nice to meet you all.")]
    chinese = [TranscriptSegment(0.0, 2.0, "很高兴认识大家。")]

    def cues(segments: list[TranscriptSegment], is_vertical: bool) -> list[str]:
        document_ = document(segments=segments, is_vertical=is_vertical)
        content = export(tmp_path / "talk", "srt", document_).read_text("utf-8")
        return [block.split("\n")[2] for block in content.split("\n\n")]

    assert cues(latin, is_vertical=False) == ["Nice to meet you all."]
    # 24 characters by default
    ConfigManager.modify_value(
        section, ConfigSubtitles.Key.MAX_LINE_WIDTH_VERTICAL, "12"
    )
    assert cues(latin, is_vertical=True) == ["Nice to meet", "you all."]
    assert cues(chinese, is_vertical=False) == ["很高兴认", "识大家。"]


def test_vtt(tmp_path: Path) -> None:
    content = export(tmp_path / "talk", "vtt", document()).read_text()

    assert content.startswith("WEBVTT\n\n00:00:00.000 --> 00:00:01.500\nHello.\n")


def test_tsv(tmp_path: Path) -> None:
    content = export(tmp_path / "talk", "tsv", document()).read_text()

    assert content.splitlines() == [
        "start\tend\ttext",
        "0\t1500\tHello.",
        "3661250\t3662000\t[SPEAKER_01]: Bye.",
    ]


def test_json_and_txt_use_the_text(tmp_path: Path) -> None:
    data = json.loads(
        export(tmp_path / "talk", "json", document(text="Edited")).read_text()
    )
    txt = export(tmp_path / "talk", "txt", document(text="Edited", segments=[]))

    assert data["text"] == "Edited"
    assert data["segments"][1]["speaker"] == "SPEAKER_01"
    assert "summary" not in data
    assert txt.read_text() == "Edited\n"


def test_json_includes_the_summary(tmp_path: Path) -> None:
    path = export(tmp_path / "talk", "json", document(summary=SUMMARY))

    assert json.loads(path.read_text())["summary"]["summary"] == "A greeting."


def test_the_summary_is_exported_to_documents(tmp_path: Path) -> None:
    # Without the formats of the subtitles or the data of the transcription
    assert list(SUMMARY_EXPORTERS) == ["txt", "md", "docx"]
    with pytest.raises(ValueError):
        export_summary(tmp_path / "talk", "srt", "Talk", SUMMARY)

    txt = export_summary(tmp_path / "talk", "txt", "Talk", SUMMARY).read_text(
        encoding="utf-8"
    )
    assert txt == (
        "Summary\n\nA greeting.\n\nKey points\n\n• They say hello\n\n"
        "Chapters\n\n00:00  Greeting\n01:01:01  Farewell\n"
    )

    md = export_summary(tmp_path / "talk", "md", "Talk", SUMMARY).read_text(
        encoding="utf-8"
    )
    assert md.startswith("# Talk\n\n## Summary\n\nA greeting.")
    assert "- **01:01:01** Farewell" in md
    # Only the summary, not the transcription
    assert "Transcription" not in md

    docx = Document(str(export_summary(tmp_path / "talk", "docx", "Talk", SUMMARY)))
    texts = [paragraph.text for paragraph in docx.paragraphs]
    assert texts[:3] == ["Talk", "Summary", "A greeting."]
    assert "Transcription" not in texts


def test_timed_formats_need_segments(tmp_path: Path) -> None:
    with pytest.raises(ValueError):
        export(tmp_path / "talk", "srt", document(segments=[]))

    assert available_formats(has_segments=False) == ["txt", "md", "docx", "json"]
    assert set(available_formats(has_segments=True)) == {
        "txt",
        "md",
        "docx",
        "srt",
        "vtt",
        "ass",
        "tsv",
        "json",
    }


def test_paragraphs_follow_the_turns_of_the_speakers() -> None:
    segments = [
        TranscriptSegment(0, 1, "Hi.", speaker="SPEAKER_00"),
        TranscriptSegment(1, 2, "How are you?", speaker="SPEAKER_00"),
        TranscriptSegment(2, 3, "Fine.", speaker="SPEAKER_01"),
    ]

    assert paragraphs(document(segments=segments)) == [
        Paragraph("Hi. How are you?", 0, "SPEAKER_00"),
        Paragraph("Fine.", 2, "SPEAKER_01"),
    ]


def test_paragraphs_without_speakers_are_split_at_the_pauses() -> None:
    segments = [
        TranscriptSegment(0, 1, "One."),
        TranscriptSegment(1.5, 2, "Two."),
        TranscriptSegment(10, 11, "Three."),
    ]

    assert paragraphs(document(segments=segments)) == [
        Paragraph("One. Two.", 0),
        Paragraph("Three.", 10),
    ]


def test_paragraphs_of_the_edited_text_are_its_blocks() -> None:
    edited = document(text="First.\n\nSecond.\n", is_text_edited=True)

    assert paragraphs(edited) == [Paragraph("First."), Paragraph("Second.")]


def test_markdown(tmp_path: Path) -> None:
    content = export(tmp_path / "talk", "md", document(summary=SUMMARY)).read_text()

    assert content.startswith("# Talk\n\n## Summary\n\nA greeting.\n")
    assert "- They say hello" in content
    assert "- **01:01:01** Farewell" in content
    assert "**[00:00]** Hello." in content
    assert "**[01:01:01] SPEAKER_01:** Bye." in content


def test_markdown_without_timestamps(tmp_path: Path) -> None:
    content = export(tmp_path / "talk", "md", document(segments=[])).read_text()

    assert content == "# Talk\n\nHello. Bye.\n"


def test_docx(tmp_path: Path) -> None:
    path = export(tmp_path / "talk", "docx", document(summary=SUMMARY))

    texts = [paragraph.text for paragraph in Document(str(path)).paragraphs]
    assert path.suffix == ".docx"
    assert texts[0] == "Talk"
    assert "A greeting." in texts
    assert "[01:01:01] SPEAKER_01: Bye." in texts


def test_several_transcriptions_are_exported_together(tmp_path: Path) -> None:
    documents = [
        document(title="a.mp3", segments=[], text="Hello."),
        document(title="b.mp3", segments=[], text="Bye."),
    ]

    txt = export_collection(tmp_path / "talks", "txt", "Talks", documents)
    md = export_collection(tmp_path / "talks", "md", "Talks", documents)
    docx = export_collection(tmp_path / "talks", "docx", "Talks", documents)

    assert txt.read_text(encoding="utf-8") == (
        "a.mp3\n=====\n\nHello.\n\n\nb.mp3\n=====\n\nBye.\n"
    )
    assert md.read_text(encoding="utf-8") == (
        "# Talks\n\n## a.mp3\n\nHello.\n\n## b.mp3\n\nBye.\n"
    )
    headings = [
        paragraph.text
        for paragraph in Document(str(docx)).paragraphs
        if paragraph.style.name.startswith(("Heading", "Title"))
    ]
    assert headings == ["Talks", "a.mp3", "b.mp3"]
    with pytest.raises(ValueError):
        export_collection(tmp_path / "talks", "srt", "Talks", documents)
