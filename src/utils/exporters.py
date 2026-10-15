"""
Exports the transcriptions of the history to files. Unlike the WhisperX writers,
they only need the stored segments (or text), so any entry of the history can be
exported at any time.
"""

import json
from collections.abc import Callable
from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path
from typing import Any

from models.summary import TranscriptSummary
from models.transcript_segment import TranscriptSegment
from utils.config_manager import ConfigManager
from utils.i18n import _
from utils.subtitle_cues import Cue, build_cues, line_limits
from utils.time_format import format_subtitle_time, format_timestamp

# A paragraph of a document is closed after a pause this long, or once it's this
# long, when the speakers are not identified
PARAGRAPH_PAUSE_SECONDS = 2.0
PARAGRAPH_MAX_CHARS = 600

# The colors of the ASS subtitles, as `&HAABBGGRR`: white without speakers, and one
# color per speaker, light enough to be read over the video
ASS_DEFAULT_COLOR = "&H00FFFFFF"
ASS_SPEAKER_COLORS = (
    "&H0000FFFF",  # Yellow
    "&H00FFFF00",  # Cyan
    "&H0080FF80",  # Green
    "&H00FF80FF",  # Pink
    "&H000080FF",  # Orange
    "&H00FFB080",  # Blue
    "&H00FF80B0",  # Purple
    "&H008080FF",  # Red
)
# For the resolution of the script (1920x1080, or 1080x1920 if vertical), which the
# players scale to the one of the video
ASS_FONT_SIZE = 64
ASS_FONT_SIZE_VERTICAL = 56
ASS_MARGIN = 60
ASS_MARGIN_VERTICAL = 160


@dataclass
class ExportDocument:
    """A transcription to export."""

    title: str
    text: str
    segments: list[TranscriptSegment]
    summary: TranscriptSummary | None = None
    # The text edited by the user is the one exported in the formats without
    # timestamps, since the segments keep the original transcription
    is_text_edited: bool = False
    # Whether the video is taller than it's wide, which narrows its subtitles
    is_vertical: bool = False


@dataclass(frozen=True)
class Paragraph:
    text: str
    # Only known if the transcription has timestamps
    start: float | None = None
    speaker: str | None = None


def _cue_text(segment: TranscriptSegment) -> str:
    return f"[{segment.speaker}]: {segment.text}" if segment.speaker else segment.text


def speaker_cues(
    segments: list[TranscriptSegment], is_vertical: bool = False
) -> list[tuple[Cue, str | None]]:
    """
    The segments split into subtitles, as the user configured them, with the
    speaker of each one.
    """
    config = ConfigManager.get_config_subtitles()
    width, count = line_limits(config, segments, is_vertical)
    cues: list[tuple[Cue, str | None]] = []
    for segment in segments:
        segment_cues = build_cues([segment], width, count)
        # The speaker is added after splitting, so the subtitles are the ones
        # shown with the transcript and over the video, which show it apart
        if segment.speaker and segment_cues:
            first = segment_cues[0]
            segment_cues[0] = replace(first, text=f"[{segment.speaker}]: {first.text}")
        cues.extend((cue, segment.speaker) for cue in segment_cues)
    return cues


def subtitle_cues(
    segments: list[TranscriptSegment], is_vertical: bool = False
) -> list[Cue]:
    """The segments split into subtitles, as the user configured them."""
    return [cue for cue, _speaker in speaker_cues(segments, is_vertical)]


def segments_to_srt(
    segments: list[TranscriptSegment], is_vertical: bool = False
) -> str:
    cues = [
        f"{idx}\n{format_subtitle_time(cue.start, ',')} --> "
        f"{format_subtitle_time(cue.end, ',')}\n{cue.text}\n"
        for idx, cue in enumerate(subtitle_cues(segments, is_vertical), start=1)
    ]
    return "\n".join(cues)


def segments_to_vtt(
    segments: list[TranscriptSegment], is_vertical: bool = False
) -> str:
    cues = [
        f"{format_subtitle_time(cue.start, '.')} --> "
        f"{format_subtitle_time(cue.end, '.')}\n{cue.text}\n"
        for cue in subtitle_cues(segments, is_vertical)
    ]
    return "WEBVTT\n\n" + "\n".join(cues)


def _ass_time(seconds: float) -> str:
    """Formats a position as `H:MM:SS.cc`, the times of the ASS subtitles."""
    centiseconds = round(max(seconds, 0) * 100)
    hours, centiseconds = divmod(centiseconds, 360_000)
    minutes, centiseconds = divmod(centiseconds, 6_000)
    seconds_part, centiseconds = divmod(centiseconds, 100)
    return f"{hours}:{minutes:02}:{seconds_part:02}.{centiseconds:02}"


def _ass_text(text: str) -> str:
    """
    Escapes the text of a subtitle: the braces start the tags of ASS, and the
    backslash its line breaks, which can't be escaped.
    """
    text = text.replace("\\", "/").replace("{", "(").replace("}", ")")
    return text.replace("\n", "\\N")


def _ass_style(name: str, color: str, is_vertical: bool) -> str:
    # Like the subtitles shown over the video: white or colored, with a black
    # outline and a shadow, centered at the bottom
    font_size = ASS_FONT_SIZE_VERTICAL if is_vertical else ASS_FONT_SIZE
    margin = ASS_MARGIN_VERTICAL if is_vertical else ASS_MARGIN
    return (
        f"Style: {name},Arial,{font_size},{color},{color},&H00000000,&H80000000,"
        f"0,0,0,0,100,100,0,0,1,3,1,2,{margin},{margin},{margin},1"
    )


def segments_to_ass(
    segments: list[TranscriptSegment], is_vertical: bool = False
) -> str:
    """
    The subtitles in the Advanced SubStation Alpha format, which video editors and
    players (e.g. VLC or mpv) show with the style of the file. Each speaker has a
    color of its own, so their lines can be told apart.
    """
    cues = speaker_cues(segments, is_vertical)
    speakers = list(dict.fromkeys(speaker for _cue, speaker in cues if speaker))
    styles = {speaker: f"Speaker{idx}" for idx, speaker in enumerate(speakers, start=1)}
    width, height = (1080, 1920) if is_vertical else (1920, 1080)

    lines = [
        "[Script Info]",
        "ScriptType: v4.00+",
        f"PlayResX: {width}",
        f"PlayResY: {height}",
        "WrapStyle: 2",
        "ScaledBorderAndShadow: yes",
        "",
        "[V4+ Styles]",
        "Format: Name, Fontname, Fontsize, PrimaryColour, SecondaryColour, "
        "OutlineColour, BackColour, Bold, Italic, Underline, StrikeOut, ScaleX, "
        "ScaleY, Spacing, Angle, BorderStyle, Outline, Shadow, Alignment, "
        "MarginL, MarginR, MarginV, Encoding",
        _ass_style("Default", ASS_DEFAULT_COLOR, is_vertical),
    ]
    lines += [
        _ass_style(
            style, ASS_SPEAKER_COLORS[idx % len(ASS_SPEAKER_COLORS)], is_vertical
        )
        for idx, style in enumerate(styles.values())
    ]
    lines += [
        "",
        "[Events]",
        "Format: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, "
        "Effect, Text",
    ]
    lines += [
        f"Dialogue: 0,{_ass_time(cue.start)},{_ass_time(cue.end)},"
        f"{styles.get(speaker or '', 'Default')},{(speaker or '').replace(',', ' ')},"
        f"0,0,0,,{_ass_text(cue.text)}"
        for cue, speaker in cues
    ]
    return "\n".join(lines) + "\n"


def paragraphs(document: ExportDocument) -> list[Paragraph]:
    """
    Splits the transcription into the paragraphs of a document: one per turn of
    each speaker, or after each pause if the speakers are not identified. The
    edited text is split by its blank lines.
    """
    if not document.segments or document.is_text_edited:
        return [
            Paragraph(block.strip())
            for block in document.text.split("\n\n")
            if block.strip()
        ]

    has_speakers = any(segment.speaker for segment in document.segments)
    result: list[Paragraph] = []
    texts: list[str] = []
    start: float | None = None
    speaker: str | None = None
    previous_end = 0.0

    def close() -> None:
        if texts:
            result.append(Paragraph(" ".join(texts), start, speaker))
            texts.clear()

    for segment in document.segments:
        if has_speakers:
            is_new = segment.speaker != speaker
        else:
            is_new = (
                segment.start - previous_end >= PARAGRAPH_PAUSE_SECONDS
                or sum(len(text) for text in texts) >= PARAGRAPH_MAX_CHARS
            )

        if is_new or not texts:
            close()
            start, speaker = segment.start, segment.speaker

        texts.append(segment.text)
        previous_end = segment.end

    close()
    return result


def to_srt(document: ExportDocument) -> str:
    return segments_to_srt(document.segments, document.is_vertical)


def to_vtt(document: ExportDocument) -> str:
    return segments_to_vtt(document.segments, document.is_vertical)


def to_ass(document: ExportDocument) -> str:
    return segments_to_ass(document.segments, document.is_vertical)


def to_tsv(document: ExportDocument) -> str:
    rows = ["start\tend\ttext"]
    rows.extend(
        f"{round(s.start * 1000)}\t{round(s.end * 1000)}\t"
        f"{_cue_text(s).replace(chr(9), ' ')}"
        for s in document.segments
    )
    return "\n".join(rows) + "\n"


def to_json(document: ExportDocument) -> str:
    data: dict[str, object] = {
        "text": document.text,
        "segments": [s.to_dict() for s in document.segments],
    }
    if document.summary:
        data["summary"] = document.summary.to_dict()
    return json.dumps(data, ensure_ascii=False, indent=2)


def to_txt(document: ExportDocument) -> str:
    text = document.text
    return text if text.endswith("\n") else f"{text}\n"


def to_markdown(document: ExportDocument) -> str:
    lines = [f"# {document.title}", ""]

    if summary := document.summary:
        lines += _markdown_summary(summary)
        lines += ["## " + _("Transcription"), ""]

    for paragraph in paragraphs(document):
        prefix = ""
        if paragraph.start is not None:
            prefix = f"[{format_timestamp(paragraph.start)}] "
        if paragraph.speaker:
            prefix += f"{paragraph.speaker}: "
        lines += [
            f"**{prefix.strip()}** {paragraph.text}" if prefix else paragraph.text,
            "",
        ]

    return "\n".join(lines).rstrip() + "\n"


def _markdown_summary(summary: TranscriptSummary) -> list[str]:
    """The lines of the summary, its key points and its chapters, in Markdown."""
    lines = ["## " + _("Summary"), "", summary.summary, ""]
    if summary.key_points:
        lines += ["### " + _("Key points"), ""]
        lines += [f"- {point}" for point in summary.key_points]
        lines.append("")
    if summary.chapters:
        lines += ["### " + _("Chapters"), ""]
        lines += [
            f"- **{format_timestamp(chapter.start)}** {chapter.title}"
            for chapter in summary.chapters
        ]
        lines.append("")
    return lines


def _new_docx(title: str) -> Any:
    """A Word document with the title."""
    # Imported here because it's only needed to export Word documents
    from docx import Document

    docx = Document()
    docx.core_properties.title = title
    docx.add_heading(title, level=0)
    return docx


def _add_docx_summary(docx: Any, summary: TranscriptSummary) -> None:
    """Adds the summary, its key points and its chapters to a Word document."""
    docx.add_heading(_("Summary"), level=1)
    docx.add_paragraph(summary.summary)
    if summary.key_points:
        docx.add_heading(_("Key points"), level=2)
        for point in summary.key_points:
            docx.add_paragraph(point, style="List Bullet")
    if summary.chapters:
        docx.add_heading(_("Chapters"), level=2)
        for chapter in summary.chapters:
            chapter_paragraph = docx.add_paragraph(style="List Bullet")
            chapter_paragraph.add_run(format_timestamp(chapter.start)).bold = True
            chapter_paragraph.add_run(f"  {chapter.title}")


def _docx_bytes(docx: Any) -> bytes:
    file = BytesIO()
    docx.save(file)
    return file.getvalue()


def to_docx(document: ExportDocument) -> bytes:
    docx = _new_docx(document.title)

    if summary := document.summary:
        _add_docx_summary(docx, summary)
        docx.add_heading(_("Transcription"), level=1)

    for paragraph in paragraphs(document):
        docx_paragraph = docx.add_paragraph()
        if paragraph.start is not None:
            docx_paragraph.add_run(
                f"[{format_timestamp(paragraph.start)}] "
            ).bold = True
        if paragraph.speaker:
            docx_paragraph.add_run(f"{paragraph.speaker}: ").bold = True
        docx_paragraph.add_run(paragraph.text)

    return _docx_bytes(docx)


Exporter = Callable[[ExportDocument], str | bytes]

EXPORTERS: dict[str, Exporter] = {
    "txt": to_txt,
    "md": to_markdown,
    "docx": to_docx,
    "srt": to_srt,
    "vtt": to_vtt,
    "ass": to_ass,
    "tsv": to_tsv,
    "json": to_json,
}
# Formats that need the timestamps of the segments
TIMED_FORMATS = {"srt", "vtt", "ass", "tsv"}


def available_formats(has_segments: bool) -> list[str]:
    """The formats a transcription can be exported to."""
    return [
        file_type
        for file_type in EXPORTERS
        if has_segments or file_type not in TIMED_FORMATS
    ]


def export(file_path: Path, file_type: str, document: ExportDocument) -> Path:
    """
    Writes the transcription to a file.

    :param file_path: The path of the file. Its extension is replaced by the one of
                      the file type.
    :param file_type: One of `EXPORTERS`.
    :raises ValueError: If the format needs segments and there are none.
    :return: The path of the written file.
    """
    if file_type not in EXPORTERS:
        raise ValueError(f"Unsupported file type: {file_type}")

    if file_type in TIMED_FORMATS and not document.segments:
        raise ValueError(f"The .{file_type} format needs the timestamps")

    return _write(file_path, file_type, EXPORTERS[file_type](document))


def summary_to_txt(_title: str, summary: TranscriptSummary) -> str:
    parts = [_("Summary"), summary.summary]
    if summary.key_points:
        parts += [
            _("Key points"),
            "\n".join(f"• {point}" for point in summary.key_points),
        ]
    if summary.chapters:
        parts += [
            _("Chapters"),
            "\n".join(
                f"{format_timestamp(chapter.start)}  {chapter.title}"
                for chapter in summary.chapters
            ),
        ]
    return "\n\n".join(parts) + "\n"


def summary_to_markdown(title: str, summary: TranscriptSummary) -> str:
    return "\n".join([f"# {title}", "", *_markdown_summary(summary)]).rstrip() + "\n"


def summary_to_docx(title: str, summary: TranscriptSummary) -> bytes:
    docx = _new_docx(title)
    _add_docx_summary(docx, summary)
    return _docx_bytes(docx)


# The summary is only exported to documents: it has no timestamps of its own to
# make subtitles, and its data is already in the JSON of the transcription
SUMMARY_EXPORTERS: dict[str, Callable[[str, TranscriptSummary], str | bytes]] = {
    "txt": summary_to_txt,
    "md": summary_to_markdown,
    "docx": summary_to_docx,
}


def export_summary(
    file_path: Path, file_type: str, title: str, summary: TranscriptSummary
) -> Path:
    """
    Writes the summary of a transcription to a file.

    :param file_type: One of `SUMMARY_EXPORTERS`.
    :raises ValueError: If the summary can't be exported to the file type.
    :return: The path of the written file.
    """
    if file_type not in SUMMARY_EXPORTERS:
        raise ValueError(f"Unsupported file type: {file_type}")
    return _write(file_path, file_type, SUMMARY_EXPORTERS[file_type](title, summary))


def _write(file_path: Path, file_type: str, content: str | bytes) -> Path:
    """Writes the content with the extension of the file type."""
    output_path = file_path.with_suffix(f".{file_type}")
    if isinstance(content, bytes):
        output_path.write_bytes(content)
    else:
        output_path.write_text(content, encoding="utf-8")
    return output_path
